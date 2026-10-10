"""Read .xlsx workbooks into classified rows. The source file is never saved."""

from __future__ import annotations

from dataclasses import dataclass, field
from pathlib import Path
from typing import Any
from zipfile import ZipFile

from openpyxl import load_workbook
from openpyxl.worksheet.worksheet import Worksheet

from app.importing.images import AssociatedImage, associate_images, read_all_anchors
from app.importing.normalize import (
    build_match_key,
    classify_price,
    collapse_whitespace,
    description_sha256,
    detect_brands,
    display_model,
    filename_hint,
    is_cloud_history_model,
    is_numeric_value,
    is_service_option,
    normalize_header,
    normalize_model,
    normalize_option,
    trim_description,
)

MAX_UPLOAD_BYTES = 40 * 1024 * 1024
_SCAN_LIMIT = 20_000
_HEADER_SCAN = 30
_KNOWN_HEADERS = frozenset(
    {"№", "no", "фото", "модель", "характеристика", "доп-опция", "цена", "цена для дилера"}
)
_MODEL = "модель"
_PHOTO = "фото"
_DESCRIPTION = "характеристика"
_OPTION = "доп-опция"
_PRICE = "цена"
_DEALER = "цена для дилера"
_LINE = "№"


class MissingCache:
    def __init__(self, formula: str) -> None:
        self.formula = formula


@dataclass
class ParsedRow:
    worksheet_name: str
    source_row: int
    classification: str
    raw_cells: dict[str, Any]
    proposal: dict[str, Any]
    messages: list[dict[str, Any]]
    action: str = "skip"
    resolution: str = "pending"
    matched_product_id: int | None = None


@dataclass
class ParsedSheet:
    name: str
    header_row: int | None
    rows: list[ParsedRow] = field(default_factory=list)
    errors: list[dict[str, Any]] = field(default_factory=list)
    warnings: list[dict[str, Any]] = field(default_factory=list)
    images: list[AssociatedImage] = field(default_factory=list)
    trailing_empty_rows: int = 0


@dataclass
class ParsedWorkbook:
    source_code: str
    sheets: list[ParsedSheet]
    errors: list[dict[str, Any]] = field(default_factory=list)
    warnings: list[dict[str, Any]] = field(default_factory=list)

    @property
    def has_required_sheet(self) -> bool:
        return any(sheet.header_row is not None for sheet in self.sheets)


@dataclass
class WorkbookInspection:
    container_errors: list[dict[str, Any]]
    signature: str | None
    signatures: list[str]
    filename_hint: str | None
    sheet_headers: list[dict[str, Any]]
    warnings: list[dict[str, Any]]


def inspect_container(path: Path, original_filename: str) -> WorkbookInspection:
    errors = container_errors(path)
    hint = filename_hint(original_filename)
    if errors:
        return WorkbookInspection(errors, None, [], hint, [], [])
    warnings: list[dict[str, Any]] = []
    if has_macros(path):
        warnings.append(_warning("macros_not_executed", "The workbook contains macros. They were not executed."))
    try:
        workbook = load_workbook(path, read_only=False, data_only=False, keep_links=False)
    except Exception:
        return WorkbookInspection(
            [_error("corrupt_workbook", "The workbook could not be read.")],
            None,
            [],
            hint,
            [],
            warnings,
        )
    try:
        sheet_headers = [sheet_header(worksheet) for worksheet in workbook.worksheets]
    finally:
        workbook.close()
    signatures = sorted({item["signature"] for item in sheet_headers if item["signature"]})
    signature = signatures[0] if len(signatures) == 1 else None
    if len(signatures) > 1:
        errors = [_error("ambiguous_signature", "The workbook matches both Hikvision and EZVIZ layouts.")]
    return WorkbookInspection(errors, signature, signatures, hint, sheet_headers, warnings)


def read_workbook(path: Path, source_code: str) -> ParsedWorkbook:
    """Parse one identified workbook. The caller has already checked the container."""

    if source_code not in {"hikvision", "ezviz"}:
        raise ValueError("source_code must be hikvision or ezviz")
    workbook = load_workbook(path, read_only=False, data_only=False, keep_links=False)
    data_only = None
    try:
        if _workbook_has_formula(workbook):
            data_only = load_workbook(path, read_only=False, data_only=True, keep_links=False)
        image_anchors, image_warnings = read_all_anchors(path)
        sheets: list[ParsedSheet] = []
        errors: list[dict[str, Any]] = []
        warnings: list[dict[str, Any]] = list(image_warnings)
        if has_macros(path):
            warnings.append(
                _warning("macros_not_executed", "The workbook contains macros. They were not executed.")
            )
        for worksheet in workbook.worksheets:
            cached_sheet = data_only[worksheet.title] if data_only is not None else None
            sheet = _parse_sheet(
                worksheet,
                cached_sheet,
                source_code,
                image_anchors.get(worksheet.title, []),
            )
            if sheet.header_row is None:
                errors.extend(sheet.errors)
            warnings.extend(sheet.warnings)
            sheets.append(sheet)
        if not any(sheet.header_row is not None for sheet in sheets):
            errors.append(_error("missing_required_header", "The workbook has no sheet with the required headers."))
        return ParsedWorkbook(source_code, sheets, errors, warnings)
    finally:
        workbook.close()
        if data_only is not None:
            data_only.close()


def container_errors(path: Path) -> list[dict[str, Any]]:
    if path.suffix.casefold() != ".xlsx":
        return [_error("unsupported_type", "Only .xlsx workbooks can be staged.")]
    try:
        size = path.stat().st_size
    except OSError:
        return [_error("unreadable_file", "The workbook could not be read.")]
    if size > MAX_UPLOAD_BYTES:
        return [_error("file_too_large", "The workbook is larger than 40 MB.")]
    try:
        with path.open("rb") as handle:
            signature = handle.read(4)
    except OSError:
        return [_error("unreadable_file", "The workbook could not be read.")]
    if signature not in {b"PK\x03\x04", b"PK\x05\x06", b"PK\x07\x08"}:
        return [_error("invalid_zip", "The workbook is not a zip package.")]
    try:
        with ZipFile(path) as package:
            names = package.namelist()
    except Exception:
        return [_error("invalid_zip", "The workbook is not a zip package.")]
    if any(_unsafe_zip_name(name) for name in names):
        return [_error("invalid_zip", "The workbook package contains an unsafe path.")]
    if "xl/workbook.xml" not in names:
        return [_error("missing_workbook_xml", "The workbook package has no xl/workbook.xml.")]
    return []


def _unsafe_zip_name(name: str) -> bool:
    normalized = name.replace("\\", "/")
    if not normalized or normalized.startswith("/") or ":" in normalized:
        return True
    return any(part == ".." for part in normalized.split("/"))


def has_macros(path: Path) -> bool:
    try:
        with ZipFile(path) as package:
            return "xl/vbaProject.bin" in package.namelist()
    except Exception:
        return False


def sheet_header(worksheet: Worksheet) -> dict[str, Any]:
    limit = min(worksheet.max_row or 1, _HEADER_SCAN)
    max_column = worksheet.max_column or 1
    for row_number in range(1, limit + 1):
        columns: dict[str, int] = {}
        warnings: list[dict[str, Any]] = []
        for column in range(1, max_column + 1):
            value = worksheet.cell(row_number, column).value
            if not isinstance(value, str) or not value.strip():
                continue
            name = normalize_header(value)
            if name in columns:
                warnings.append(
                    _warning(
                        "duplicate_header",
                        f"Repeated header {value.strip()} uses the first column.",
                        worksheet.title,
                    )
                )
                continue
            columns[name] = column
            if name not in _KNOWN_HEADERS:
                warnings.append(
                    _warning(
                        "unmapped_column",
                        f"Column {value.strip()} is not mapped.",
                        worksheet.title,
                    )
                )
        signature = header_signature(set(columns))
        if signature is not None:
            return {
                "worksheet": worksheet.title,
                "header_row": row_number,
                "columns": columns,
                "signature": signature,
                "warnings": warnings,
            }
    return {
        "worksheet": worksheet.title,
        "header_row": None,
        "columns": {},
        "signature": None,
        "warnings": [],
    }


def header_signature(headers: set[str]) -> str | None:
    if _MODEL in headers and _DEALER in headers:
        return "ezviz"
    if _MODEL in headers and _PRICE in headers and _DEALER not in headers:
        return "hikvision"
    return None


def _parse_sheet(
    worksheet: Worksheet,
    cached_sheet: Worksheet | None,
    source_code: str,
    anchors: list[Any],
) -> ParsedSheet:
    header = sheet_header(worksheet)
    if header["signature"] != source_code:
        message = (
            "The sheet is missing Модель or Цена для дилера."
            if source_code == "ezviz"
            else "The sheet is missing Модель or Цена."
        )
        return ParsedSheet(
            name=worksheet.title,
            header_row=None,
            errors=[_error("missing_required_header", message, worksheet.title)],
        )
    header_row = int(header["header_row"])
    columns: dict[str, int] = header["columns"]
    note_column = _note_column(worksheet, header_row, columns) if source_code == "hikvision" else None
    needed = list(columns.values())
    if note_column is not None:
        needed.append(note_column)
    grid = _read_grid(worksheet, cached_sheet, header_row, needed)
    last_row = header_row
    for row_number, _column in grid:
        last_row = max(last_row, row_number)
    for anchor in anchors:
        if anchor.row is not None:
            last_row = max(last_row, anchor.row)
    declared_last = worksheet.max_row or header_row
    rows = [
        _classify_row(
            worksheet.title,
            row_number,
            columns,
            note_column,
            grid,
            source_code,
            header,
        )
        for row_number in range(header_row + 1, last_row + 1)
    ]
    use_headings = source_code == "hikvision" and any(row.classification == "heading" for row in rows)
    _assign_categories(rows, use_headings=use_headings, sheet_name=worksheet.title)
    model_rows = {row.source_row for row in rows if row.classification in {"product", "service"}}
    merges = [
        (merged.min_row, merged.min_col, merged.max_row, merged.max_col)
        for merged in worksheet.merged_cells.ranges
    ]
    images = associate_images(
        anchors,
        model_rows=model_rows,
        photo_column=columns.get(_PHOTO),
        photo_merges=merges,
    )
    _attach_images(rows, images)
    return ParsedSheet(
        name=worksheet.title,
        header_row=header_row,
        rows=rows,
        warnings=list(header["warnings"]),
        images=images,
        trailing_empty_rows=max(0, declared_last - last_row),
    )


def _note_column(worksheet: Worksheet, header_row: int, columns: dict[str, int]) -> int | None:
    price_column = columns.get(_PRICE)
    if price_column is None:
        return None
    candidate = price_column + 1
    header = worksheet.cell(header_row, candidate).value
    if isinstance(header, str) and header.strip():
        return None
    return candidate


def _read_grid(
    worksheet: Worksheet,
    cached_sheet: Worksheet | None,
    header_row: int,
    columns: list[int],
) -> dict[tuple[int, int], Any]:
    grid: dict[tuple[int, int], Any] = {}
    last = min(worksheet.max_row or header_row, header_row + _SCAN_LIMIT)
    for row_number in range(header_row + 1, last + 1):
        for column in columns:
            cell = worksheet.cell(row_number, column)
            if cell.data_type == "f":
                cached = cached_sheet.cell(row_number, column).value if cached_sheet is not None else None
                if cached is None or (isinstance(cached, str) and cached.strip() == ""):
                    grid[(row_number, column)] = MissingCache(str(cell.value))
                else:
                    grid[(row_number, column)] = cached
                continue
            if cell.value is None:
                continue
            if isinstance(cell.value, str) and cell.value.strip() == "":
                continue
            grid[(row_number, column)] = cell.value
    return grid


def _classify_row(
    worksheet_name: str,
    row_number: int,
    columns: dict[str, int],
    note_column: int | None,
    grid: dict[tuple[int, int], Any],
    source_code: str,
    header: dict[str, Any],
) -> ParsedRow:
    def take(name: str) -> Any:
        column = columns.get(name)
        if column is None:
            return None
        return grid.get((row_number, column))

    messages: list[dict[str, Any]] = []
    line_value = take(_LINE) if _LINE in columns else take("no")
    model_value = take(_MODEL)
    description_value = take(_DESCRIPTION)
    option_value = take(_OPTION)
    price_header = _DEALER if source_code == "ezviz" else _PRICE
    price_value = take(price_header)
    note_value = grid.get((row_number, note_column)) if note_column is not None else None

    raw_cells: dict[str, Any] = {
        "№": _raw(line_value),
        "Модель": _raw(model_value),
        "Характеристика": _raw(description_value),
        "Доп-опция": _raw(option_value) if source_code == "ezviz" else None,
        "price": _raw(price_value),
        "internal_note": _raw(note_value),
    }
    if any(
        isinstance(item, MissingCache)
        for item in (model_value, description_value, option_value, price_value, note_value)
    ):
        messages.append(_message("formula_without_cache"))

    model_text = _text(model_value)
    description_text = _text(description_value)
    option_text = _text(option_value) if source_code == "ezviz" else None
    note_text = _text(note_value)
    price = classify_price(None if isinstance(price_value, MissingCache) else price_value, missing_cache=isinstance(price_value, MissingCache))
    price["header"] = "Цена для дилера" if source_code == "ezviz" else "Цена"

    empty_mapped = (
        model_text is None
        and description_text is None
        and option_text is None
        and price["kind"] == "blank"
        and not isinstance(price_value, MissingCache)
        and note_text is None
    )
    line_only = (
        model_text is None
        and description_text is None
        and option_text is None
        and price["kind"] == "blank"
        and not isinstance(price_value, MissingCache)
        and note_text is None
        and is_numeric_value(line_value)
    )
    if isinstance(model_value, MissingCache) or isinstance(price_value, MissingCache):
        classification = "invalid"
    elif model_text is not None:
        if price["kind"] == "invalid":
            classification = "invalid"
            if price["issue"] and price["issue"] != "formula_without_cache":
                messages.append(_message(str(price["issue"])))
        elif source_code == "ezviz" and (
            is_service_option(option_text) or is_cloud_history_model(model_text)
        ):
            classification = "service"
        else:
            classification = "product"
    elif line_only:
        classification = "stray"
    elif empty_mapped and _heading_text(line_value) is not None:
        classification = "heading"
    elif empty_mapped:
        classification = "blank"
    else:
        classification = "invalid"
        messages.append(_message("missing_model"))

    proposal = _proposal(
        source_code,
        worksheet_name,
        row_number,
        classification,
        model_value if isinstance(model_value, str) else model_text,
        description_text,
        option_text,
        note_text,
        price,
        line_value,
    )
    if classification == "heading":
        proposal["heading_text"] = _heading_text(line_value)
    if classification in {"product", "service"} and source_code == "hikvision":
        detected = detect_brands(model_text, description_text)
        if len(detected) > 1:
            messages.append(_message("brand_ambiguous"))
        elif not detected:
            messages.append(_message("brand_unverified"))
    if note_text and "скид" in note_text.casefold():
        messages.append(_message("internal_discount_note"))
    return ParsedRow(
        worksheet_name,
        row_number,
        classification,
        raw_cells,
        proposal,
        _dedupe_messages(messages),
    )


def _proposal(
    source_code: str,
    worksheet_name: str,
    row_number: int,
    classification: str,
    model_raw: str | None,
    description: str | None,
    option: str | None,
    note: str | None,
    price: dict[str, Any],
    line_value: Any,
) -> dict[str, Any]:
    if classification not in {"product", "service"}:
        return {"source_row": row_number, "worksheet": worksheet_name}
    display = display_model(model_raw or "")
    normalized = normalize_model(display)
    description_clean = trim_description(description)
    option_clean = option.strip() if option else None
    option_clean = option_clean or None
    digest = description_sha256(description_clean)
    option_key = normalize_option(option_clean)
    brands = ["EZVIZ"] if source_code == "ezviz" else detect_brands(model_raw, description_clean)
    brand = brands[0] if len(brands) == 1 else None
    proposal: dict[str, Any] = {
        "source_code": source_code,
        "worksheet": worksheet_name,
        "source_row": row_number,
        "line_number": _line_text(line_value),
        "model_raw": model_raw,
        "model_display": display,
        "model_normalized": normalized,
        "option_label": option_clean,
        "option_normalized": option_key,
        "description": description_clean,
        "description_sha256": digest,
        "product_kind": "service" if classification == "service" else "product",
        "price": price,
        "internal_note": note,
        "brand": brand,
        "category_name": None,
        "section_label": None,
        "parent_section_label": None,
        "match_key": build_match_key(source_code, worksheet_name, normalized, option_key, digest),
        "initial_stock_status": "in_stock",
        "translations": {"ru": {"description": description_clean}},
        "images": [],
    }
    return proposal


def _assign_categories(rows: list[ParsedRow], *, use_headings: bool, sheet_name: str) -> None:
    if not use_headings:
        for row in rows:
            if row.classification in {"product", "service"}:
                row.proposal["category_name"] = sheet_name
        return
    heading_indexes = [index for index, row in enumerate(rows) if row.classification == "heading"]
    parent_name: str | None = None
    claimed: set[int] = set()
    for position, index in enumerate(heading_indexes):
        next_index = heading_indexes[position + 1] if position + 1 < len(heading_indexes) else len(rows)
        has_products = any(
            rows[cursor].classification in {"product", "service"} for cursor in range(index + 1, next_index)
        )
        heading = rows[index]
        name = str(heading.proposal.get("heading_text") or "")
        if not has_products:
            parent_name = name
            heading.proposal["category_proposal"] = {"name": name, "role": "parent", "parent_name": None}
            continue
        heading.proposal["category_proposal"] = {
            "name": name,
            "role": "section",
            "parent_name": parent_name,
        }
        for cursor in range(index + 1, next_index):
            row = rows[cursor]
            if row.classification not in {"product", "service"}:
                continue
            row.proposal["category_name"] = name
            row.proposal["section_label"] = name
            row.proposal["parent_section_label"] = parent_name
            claimed.add(cursor)
    for index, row in enumerate(rows):
        if row.classification in {"product", "service"} and index not in claimed:
            row.proposal["category_name"] = sheet_name
            row.messages.append(_message("category_before_heading"))
            row.messages = _dedupe_messages(row.messages)


def _attach_images(rows: list[ParsedRow], images: list[AssociatedImage]) -> None:
    by_row = {row.source_row: row for row in rows}
    grouped: dict[int, list[AssociatedImage]] = {}
    for image in images:
        if image.rejected:
            if image.row is not None and image.row in by_row:
                by_row[image.row].messages.append(_message("image_rejected"))
                by_row[image.row].messages = _dedupe_messages(by_row[image.row].messages)
            continue
        if image.link_status == "unassigned" or image.row is None:
            continue
        grouped.setdefault(image.row, []).append(image)
    for row_number, attached in grouped.items():
        row = by_row.get(row_number)
        if row is None or row.classification not in {"product", "service"}:
            continue
        ordered = sorted(attached, key=lambda image: (image.index, image.link_status or ""))
        primary_assigned = False
        rendered: list[dict[str, Any]] = []
        for image in ordered:
            proposed_primary = image.link_status == "linked_high" and not primary_assigned
            if proposed_primary:
                primary_assigned = True
            rendered.append(
                {
                    "sha256": image.sha256,
                    "link_status": image.link_status,
                    "object_key": f"objects/{image.sha256[:2]}/{image.sha256}.{image.extension}",
                    "proposed_primary": proposed_primary,
                    "image_index": image.index,
                    "shared_from_row": image.shared_from_row,
                }
            )
        row.proposal["images"] = rendered


def _workbook_has_formula(workbook: Any) -> bool:
    for worksheet in workbook.worksheets:
        max_row = min(worksheet.max_row or 1, _SCAN_LIMIT)
        max_column = min(worksheet.max_column or 1, 64)
        for row in worksheet.iter_rows(min_row=1, max_row=max_row, max_col=max_column):
            for cell in row:
                if cell.data_type == "f":
                    return True
    return False


def _heading_text(value: Any) -> str | None:
    if isinstance(value, MissingCache) or is_numeric_value(value):
        return None
    text = _text(value)
    if text is None:
        return None
    return collapse_whitespace(text)


def _text(value: Any) -> str | None:
    if value is None or isinstance(value, MissingCache) or isinstance(value, bool):
        return None
    if isinstance(value, str):
        return value if value.strip() else None
    if isinstance(value, (int, float)):
        return str(value)
    rendered = str(value).strip()
    return rendered or None


def _line_text(value: Any) -> str | None:
    if isinstance(value, MissingCache) or value is None:
        return None
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    if isinstance(value, int) and not isinstance(value, bool):
        return str(value)
    text = _text(value)
    return text.strip() if text else None


def _raw(value: Any) -> Any:
    if isinstance(value, MissingCache):
        return {"formula": value.formula, "cached": None}
    if value is None or isinstance(value, bool):
        return None
    if isinstance(value, str):
        return value
    if isinstance(value, int):
        return str(value)
    if isinstance(value, float):
        if value.is_integer():
            return str(int(value))
        return str(value)
    return str(value)


def _missing(value: Any) -> bool:
    return isinstance(value, MissingCache)


def _message(code: str) -> dict[str, str]:
    return {"code": code}


def _dedupe_messages(messages: list[dict[str, Any]]) -> list[dict[str, Any]]:
    seen: set[str] = set()
    ordered: list[dict[str, Any]] = []
    for message in sorted(messages, key=lambda item: str(item.get("code"))):
        code = str(message.get("code"))
        if code in seen:
            continue
        seen.add(code)
        ordered.append(message)
    return ordered


def _error(code: str, message: str, worksheet: str | None = None) -> dict[str, Any]:
    payload: dict[str, Any] = {"level": "error", "code": code, "message": message}
    if worksheet is not None:
        payload["worksheet"] = worksheet
    return payload


def _warning(code: str, message: str, worksheet: str | None = None) -> dict[str, Any]:
    payload: dict[str, Any] = {"level": "warning", "code": code, "message": message}
    if worksheet is not None:
        payload["worksheet"] = worksheet
    return payload
