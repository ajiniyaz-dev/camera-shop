"""Small .xlsx fixtures. These are not the client's source workbooks."""

from __future__ import annotations

import struct
import zlib
from io import BytesIO
from pathlib import Path

from openpyxl import Workbook
from openpyxl.drawing.image import Image
from openpyxl.drawing.spreadsheet_drawing import AnchorMarker, OneCellAnchor
from openpyxl.drawing.xdr import XDRPositiveSize2D
from openpyxl.utils.units import pixels_to_EMU
from openpyxl.worksheet.worksheet import Worksheet


def png_bytes(red: int) -> bytes:
    def chunk(tag: bytes, data: bytes) -> bytes:
        crc = zlib.crc32(tag + data) & 0xFFFFFFFF
        return struct.pack(">I", len(data)) + tag + data + struct.pack(">I", crc)

    header = struct.pack(">IIBBBBB", 1, 1, 8, 2, 0, 0, 0)
    pixels = b"\x00" + bytes((red, 10, 20))
    return (
        b"\x89PNG\r\n\x1a\n"
        + chunk(b"IHDR", header)
        + chunk(b"IDAT", zlib.compress(pixels))
        + chunk(b"IEND", b"")
    )


def add_png(worksheet: Worksheet, row: int, column: int, payload: bytes) -> None:
    image = Image(BytesIO(payload))
    image.width = 8
    image.height = 8
    marker = AnchorMarker(col=column - 1, colOff=0, row=row - 1, rowOff=0)
    image.anchor = OneCellAnchor(
        _from=marker,
        ext=XDRPositiveSize2D(pixels_to_EMU(8), pixels_to_EMU(8)),
    )
    worksheet.add_image(image)


def hikvision_workbook(path: Path) -> None:
    workbook = Workbook()
    hilook = workbook.active
    hilook.title = "Hilook"
    hilook.append(["№", "Фото", "Модель", "Характеристика", "Цена"])
    hilook.append(["Turbo HD камера"])
    hilook.append(["ColorVu"])
    hilook.append([1, None, "HiLook IPC-1", "уличная камера", 115])
    hilook.cell(4, 6).value = "текущая цена со скидкой 60$"
    hilook.append([None, None, None, None, None])
    hilook.append([2, None, "plain cam", "без бренда", None])
    hilook.append([3, None, "DS-1", "проект", " По запросу "])
    hilook.append([4, None, "HiLook Hikvision combo", "оба имени", 10])
    add_png(hilook, 4, 2, png_bytes(1))
    add_png(hilook, 4, 2, png_bytes(2))

    cameras = workbook.create_sheet("IP камеры")
    cameras.append(["№", "Фото", "Модель", "Характеристика", "Цена"])
    cameras.append([1])
    cameras.append([2, None, "BNC", "Разъем BNC полный", 0.15])
    cameras.append([3, None, "BNC", "Разъем BNC", 0.15])
    cameras.append([4, None, "DS80HKVS\u2010VX1", "диск", 18])

    formulas = workbook.create_sheet("PTZ")
    formulas.append(["№", "Фото", "Модель", "Характеристика", "Цена"])
    formulas["C2"] = "DS-FORM"
    formulas["E2"] = "=1+1"

    notes = workbook.create_sheet("Notes")
    notes.append(["Комментарий"])
    notes.append(["не товар"])
    workbook.save(path)
    workbook.close()


def ezviz_workbook(path: Path) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Ezviz"
    sheet["A1"] = "Прайс"
    sheet.merge_cells("A1:F1")
    sheet.append(["№", "Фото", "Модель", "Характеристика", "Доп-опция", "Цена для дилера"])
    sheet.append([1, None, "CS-H8C", "описание один", None, 50])
    sheet.append([2, None, "CS-H8C", "описание два", None, 50])
    sheet.append([3, None, "7-дневная видеоистория событий", None, "Ежемесячно:", 3])
    sheet.append([4, None, "7-дневная видеоистория событий", None, "Ежегодно:", 30])
    sheet.append([5, None, "CS-DP2C", "звонок", "с солнечной панель", 92.5])
    sheet.merge_cells("B5:B6")
    add_png(sheet, 5, 2, png_bytes(9))
    workbook.save(path)
    workbook.close()


def workbook_with_duplicate_model(path: Path) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "NVR"
    sheet.append(["№", "Фото", "Модель", "Характеристика", "Цена"])
    sheet.append([1, None, "Hikvision NVR-1", "один и тот же текст", 40])
    sheet.append([2, None, "Hikvision NVR-1", "один и тот же текст", 41])
    workbook.save(path)
    workbook.close()


def single_hikvision_row(path: Path, *, model: str = "Hikvision NVR-1", price: int = 40) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "NVR"
    sheet.append(["№", "Фото", "Модель", "Характеристика", "Цена"])
    sheet.append([1, None, model, "регистратор", price])
    workbook.save(path)
    workbook.close()
